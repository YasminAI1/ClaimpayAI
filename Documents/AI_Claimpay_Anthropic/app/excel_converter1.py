# app/excel_converter.py

from typing import List, Dict, Any, Optional, Union
import pandas as pd
import logging
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment
from datetime import datetime
import os
import re

class AnswerToExcel:
    def __init__(self):
        self.excel_data = []

    def export_to_excel(self, filepath: str) -> bool:
        try:
            if not self.excel_data:
                return False
                
            # Convert to DataFrame
            df = pd.DataFrame(self.excel_data)
            
            # Create a new workbook and select the active sheet
            wb = Workbook()
            ws = wb.active
            ws.title = 'Query Results'
            
            # Write headers
            headers = list(df.columns)
            for col, header in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col)
                cell.value = header
                cell.font = Font(bold=True)
                cell.fill = PatternFill(start_color="D3D3D3", end_color="D3D3D3", fill_type="solid")
            
            # Write data
            for row_idx, row in enumerate(df.values, 2):
                for col_idx, value in enumerate(row, 1):
                    ws.cell(row=row_idx, column=col_idx, value=value)
            
            # Adjust column widths
            for column in ws.columns:
                max_length = 0
                column = [cell for cell in column]
                for cell in column:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except:
                        pass
                adjusted_width = (max_length + 2)
                ws.column_dimensions[column[0].column_letter].width = adjusted_width
            
            # Save the workbook
            wb.save(filepath)
            return True
            
        except Exception as e:
            print(f"Error exporting to Excel: {str(e)}")
            return False

    def parse_answer(self, answer_dict: dict) -> None:
        try:
            if not answer_dict.get("answer"):
                return
                
            # Split answer into lines and extract key-value pairs
            lines = answer_dict["answer"].split("\n")
            data_entry = {
                "recordid": answer_dict.get("recordid"),
                "link": answer_dict.get("link")
            }
            
            for line in lines:
                if ":" in line:
                    line = line.strip().lstrip("- ")
                    key, value = line.split(":", 1)
                    data_entry[key.strip()] = value.strip()
                    
            self.excel_data.append(data_entry)
            
        except Exception as e:
            print(f"Error parsing answer: {str(e)}")

    def clear_data(self) -> None:
        self.excel_data = []

# Usage example:
"""
# Initialize converter
excel_converter = AnswerToExcel()

# Add answers to converter
answer1 = {
    "answer": "The most recent claim in the system is:\n- Claim ID: CL123\n- Total Bill Amount: $1,234.56",
    "recordid": 123,
    "link": "http://example.com/claim/123"
}
excel_converter.parse_answer(answer1)

# Export to Excel
excel_converter.export_to_excel("claims_report.xlsx")

# Clear data if needed
excel_converter.clear_data()
"""