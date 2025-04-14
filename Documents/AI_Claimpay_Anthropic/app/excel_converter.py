# app/excel_converter.py
import pandas as pd
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import logging
from datetime import datetime, timedelta
import os
import re
from typing import List, Dict, Any, Optional, Union

class ReportQueryBuilder:
    """Builds SQL queries for complex report requests"""
    
    def __init__(self, db_manager):
        self.db_manager = db_manager
    
    def parse_report_request(self, question: str) -> Dict[str, Any]:
        """Extract report parameters from user question"""
        # Initialize defaults
        params = {
            'provider': 'Provider 1',  # Default provider
            'date_of_loss_before': None,
            'investor': None,
            'funded_last_month': False,
            'report_type': 'provider_claims'
        }
        
        # Extract provider
        provider_match = re.search(r'Provider\s+(\w+(?:\s+\w+)*)', question, re.IGNORECASE)
        if provider_match:
            params['provider'] = provider_match.group(1)
        
        # Extract date of loss filter
        date_match = re.search(r'before\s+(\d{1,2}/\d{1,2}/\d{4})', question, re.IGNORECASE)
        if date_match:
            date_str = date_match.group(1)
            try:
                # Convert MM/DD/YYYY to YYYY-MM-DD for SQL
                month, day, year = date_str.split('/')
                params['date_of_loss_before'] = f"{year}-{month.zfill(2)}-{day.zfill(2)}"
            except ValueError:
                logging.warning(f"Invalid date format: {date_str}")
        
        # Extract investor
        investor_match = re.search(r'Investor\s+(\w+(?:\s+\w+)*)', question, re.IGNORECASE)
        if investor_match:
            params['investor'] = investor_match.group(1)
        
        # Check for "last month" filter
        if 'last month' in question.lower() or 'previous month' in question.lower():
            params['funded_last_month'] = True
        
        return params
    
    def build_report_query(self, params: Dict[str, Any]) -> str:
        """Build SQL query based on extracted parameters"""
        # Base query for all purchased claims by provider
        query = """
        SELECT 
            c.recordid, 
            c.claim_id, 
            c.claim_number, 
            c.insured, 
            c.provider, 
            c.date_of_loss, 
            c.total_bill_amount,
            c.onboarding_status, 
            c.type_of_claim, 
            pp.purchase_date, 
            pp.funded_date, 
            p.investor,
            p.portfolio_id
        FROM recordsdet_claims c
        LEFT JOIN recordsdet_portfolio_purchases pp ON c.portfolio = pp.portfolio
        LEFT JOIN recordsdet_portfolios p ON c.portfolio = p.portfolio_id
        WHERE c.provider LIKE '%{0}%'
        AND c.onboarding_status = 'Purchased'
        """.format(params['provider'])
        
        # Add date of loss filter if provided
        if params['date_of_loss_before']:
            query += f" AND c.date_of_loss < '{params['date_of_loss_before']}'"
        
        # Add investor filter if provided
        if params['investor']:
            query += f" AND p.investor LIKE '%{params['investor']}%'"
        
        # Add "funded last month" filter if requested
        if params['funded_last_month']:
            # Calculate first and last day of previous month
            today = datetime.today()
            first_day = (today.replace(day=1) - timedelta(days=1)).replace(day=1)
            last_day = today.replace(day=1) - timedelta(days=1)
            
            first_day_str = first_day.strftime('%Y-%m-%d')
            last_day_str = last_day.strftime('%Y-%m-%d')
            
            query += f" AND pp.funded_date BETWEEN '{first_day_str}' AND '{last_day_str}'"
        
        # Order by date of loss descending
        query += " ORDER BY c.date_of_loss DESC"
        
        return query
    
    def execute_report_query(self, params: Dict[str, Any]) -> pd.DataFrame:
        """Execute the report query and return a DataFrame"""
        query = self.build_report_query(params)
        logging.info(f"Executing report query: {query}")
        
        # Execute query and return DataFrame
        return self.db_manager.fetch_data(query)


class AnswerToExcel:
    def __init__(self, db_manager=None):
        self.excel_data = []
        self.db_manager = db_manager
        self.query_builder = ReportQueryBuilder(db_manager) if db_manager else None
        self.excel_dir = "excel_exports"
        os.makedirs(self.excel_dir, exist_ok=True)
    
    def export_to_excel(self, filepath: str, title: str = "Query Results") -> bool:
        try:
            if not self.excel_data:
                logging.warning("No data to export")
                return False
                
            # Convert to DataFrame
            df = pd.DataFrame(self.excel_data)
            
            # Create a new workbook and select the active sheet
            wb = Workbook()
            ws = wb.active
            ws.title = title[:31]  # Excel worksheet name limitation
            
            # Add title and timestamp
            ws.merge_cells('A1:K1')
            title_cell = ws['A1']
            title_cell.value = title
            title_cell.font = Font(bold=True, size=14)
            title_cell.alignment = Alignment(horizontal='center')
            
            ws.merge_cells('A2:K2')
            timestamp_cell = ws['A2']
            timestamp_cell.value = f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
            timestamp_cell.font = Font(italic=True)
            timestamp_cell.alignment = Alignment(horizontal='center')
            
            # Configure styles
            header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
            header_font = Font(bold=True, color="FFFFFF")
            border = Border(
                left=Side(style='thin'),
                right=Side(style='thin'),
                top=Side(style='thin'),
                bottom=Side(style='thin')
            )
            
            # Write headers - starting at row 4 to leave space for title
            headers = list(df.columns)
            for col, header in enumerate(headers, 1):
                cell = ws.cell(row=4, column=col)
                cell.value = header.upper().replace('_', ' ')
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal='center')
                cell.border = border
            
            # Format currency columns
            currency_columns = ['total_bill_amount', 'total_collections', 'purchase_value']
            
            # Write data
            for row_idx, row in enumerate(df.values, 5):
                for col_idx, value in enumerate(row, 1):
                    cell = ws.cell(row=row_idx, column=col_idx)
                    
                    # Format dates
                    if isinstance(value, (pd.Timestamp, datetime)):
                        cell.value = value.strftime('%Y-%m-%d')
                        cell.alignment = Alignment(horizontal='center')
                    # Format currency
                    elif headers[col_idx-1] in currency_columns:
                        try:
                            cell.value = float(value) if not pd.isna(value) else 0
                            cell.number_format = '_($* #,##0.00_);_($* (#,##0.00);_($* "-"??_);_(@_)'
                        except (ValueError, TypeError):
                            cell.value = value
                    else:
                        cell.value = value
                    
                    # Add borders
                    cell.border = border
                    
                    # Add alternating row shading
                    if row_idx % 2 == 0:
                        cell.fill = PatternFill(start_color="E6F0FF", end_color="E6F0FF", fill_type="solid")
            
            # Adjust column widths
            for column in ws.columns:
                max_length = 0
                column_letter = column[0].column_letter
                
                for cell in column:
                    try:
                        cell_length = len(str(cell.value))
                        if cell_length > max_length:
                            max_length = cell_length
                    except:
                        pass
                
                adjusted_width = max(max_length + 2, 12)  # Min width of 12
                ws.column_dimensions[column_letter].width = adjusted_width
            
            # Add totals row for numeric columns
            if not df.empty:
                total_row = len(df) + 5
                
                # Add "TOTAL" label
                ws.cell(row=total_row, column=1).value = "TOTAL"
                ws.cell(row=total_row, column=1).font = Font(bold=True)
                
                # Add sums for currency columns
                for col_name in currency_columns:
                    if col_name in headers:
                        col_idx = headers.index(col_name) + 1
                        cell = ws.cell(row=total_row, column=col_idx)
                        start_cell = f"{get_column_letter(col_idx)}5"
                        end_cell = f"{get_column_letter(col_idx)}{total_row-1}"
                        cell.value = f"=SUM({start_cell}:{end_cell})"
                        cell.font = Font(bold=True)
                        cell.number_format = '_($* #,##0.00_);_($* (#,##0.00);_($* "-"??_);_(@_)'
                        cell.border = Border(top=Side(style='double'))
            
            # Save the workbook
            wb.save(filepath)
            logging.info(f"Successfully exported Excel file to {filepath}")
            return True
            
        except Exception as e:
            logging.error(f"Error exporting to Excel: {str(e)}")
            return False

    def parse_answer(self, answer_dict: dict) -> None:
        try:
            if not answer_dict.get("answer"):
                return
                
            # Split answer into lines and extract key-value pairs
            lines = answer_dict["answer"].split("\n")
            data_entry = {
                "recordid": answer_dict.get("recordid"),
                "link": answer_dict.get("link")
            }
            
            for line in lines:
                if ":" in line:
                    line = line.strip().lstrip("- ")
                    key, value = line.split(":", 1)
                    data_entry[key.strip()] = value.strip()
                    
            self.excel_data.append(data_entry)
            
        except Exception as e:
            logging.error(f"Error parsing answer: {str(e)}")

    def clear_data(self) -> None:
        self.excel_data = []
    
    def generate_report_filename(self, report_type: str = "report") -> str:
        """Generate a unique filename for the report"""
        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        return f"{report_type}_{timestamp}.xlsx"
    
    def generate_report(self, question: str) -> Dict[str, Any]:
        """Generate a report based on user question"""
        if not self.db_manager or not self.query_builder:
            return {
                "answer": "Report generation not available - database manager not configured",
                "is_excel": False
            }
        
        try:
            # Parse the report request to extract parameters
            params = self.query_builder.parse_report_request(question)
            
            # Execute the query to get data
            df = self.query_builder.execute_report_query(params)
            
            if df.empty:
                return {
                    "answer": f"No records found for {params['provider']} with the specified criteria.",
                    "is_excel": False
                }
            
            # Store the DataFrame data for export
            self.excel_data = df.to_dict('records')
            
            # Create report title
            title = f"Claims Report - {params['provider']}"
            if params['date_of_loss_before']:
                title += f" (DoL before {params['date_of_loss_before']})"
            if params['investor']:
                title += f" funded by {params['investor']}"
            if params['funded_last_month']:
                title += " funded last month"
            
            # Generate unique filename
            filename = self.generate_report_filename("claims_report")
            filepath = os.path.join(self.excel_dir, filename)
            
            # Export to Excel
            self.export_to_excel(filepath, title=title)
            
            # Return the result
            return {
                "answer": f"Generated report for {params['provider']} with {len(df)} records.",
                "is_excel": True,
                "excel_path": filename,
                "data": self.excel_data
            }
        
        except Exception as e:
            logging.error(f"Error generating report: {str(e)}", exc_info=True)
            return {
                "answer": f"An error occurred generating the report: {str(e)}",
                "is_excel": False
            }